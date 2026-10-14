from pathlib import Path
from datetime import date
from openpyxl import load_workbook
from braevity.core.models import Job
from braevity.core.exporter import export_jobs


def _job(title="SDE", company="Acme", url="https://x.com/1"):
    return Job.from_fields(
        title=title, company=company, location="NYC",
        url=url, posted_date="2026-10-06", source="greenhouse",
    )


def test_export_creates_file(tmp_path):
    jobs = [_job()]
    path = export_jobs(jobs, str(tmp_path))
    assert path.exists()
    assert date.today().isoformat() in path.name


def test_export_writes_header_and_row(tmp_path):
    jobs = [_job()]
    path = export_jobs(jobs, str(tmp_path))
    wb = load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "Title" in headers
    assert "Company" in headers
    assert "URL" in headers
    assert ws.cell(row=2, column=1).value == "SDE"


def test_export_appends_to_existing(tmp_path):
    export_jobs([_job(url="https://x.com/1")], str(tmp_path))
    export_jobs([_job(url="https://x.com/2")], str(tmp_path))
    path = tmp_path / f"braevity_{date.today().isoformat()}.xlsx"
    wb = load_workbook(path)
    ws = wb.active
    assert ws.max_row == 3


def test_export_empty_jobs_no_crash(tmp_path):
    path = export_jobs([], str(tmp_path))
    assert path is None
