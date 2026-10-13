import logging
from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from braevity.core.models import Job

logger = logging.getLogger(__name__)

HEADERS = ["Title", "Company", "Location", "Posted Date", "URL", "Source", "Found At"]


def export_jobs(jobs: list[Job], excel_dir: str) -> Path | None:
    if not jobs:
        return None

    dir_path = Path(excel_dir)
    dir_path.mkdir(parents=True, exist_ok=True)
    file_path = dir_path / f"braevity_{date.today().isoformat()}.xlsx"

    try:
        if file_path.exists():
            wb = load_workbook(file_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Jobs"
            ws.append(HEADERS)
    except PermissionError:
        logger.warning("Excel file locked, writing new file")
        wb = Workbook()
        ws = wb.active
        ws.title = "Jobs"
        ws.append(HEADERS)

    # Append job rows
    for job in jobs:
        ws.append([
            job.title,
            job.company,
            job.location or "",
            job.posted_date or "",
            job.url,
            job.source_ats,
            job.found_at,
        ])

    # Try save, fallback to temp file
    try:
        wb.save(file_path)
        return file_path
    except PermissionError:
        temp_path = file_path.with_suffix(".tmp.xlsx")
        logger.warning("Excel file locked on save, writing to %s", temp_path)
        wb.save(temp_path)
        return temp_path